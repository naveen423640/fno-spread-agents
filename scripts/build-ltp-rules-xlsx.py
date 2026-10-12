# Builds docs/ltp-trading-rules.xlsx from the verified practice-tool data and the knowledge graph:
#   python3 scripts/build-ltp-rules-xlsx.py   (needs openpyxl; the lookup formulas are calculated when the file is opened)
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

tool = json.load(open('docs/data/ltp-coa-practice-tool.json'))
graph = json.load(open('docs/data/ltp-knowledge-graph.json'))
nodes = {n['id']: n for n in graph['nodes']}
WORD = {'S': 'stable strong', 'B': 'bearish', 'U': 'bullish', 'HB': 'highly bearish', 'HU': 'highly bullish'}
NAME = {'strong': 'Strong', 'wtb': 'WTB (weak toward bottom)', 'wtt': 'WTT (weak toward top)'}
F = 'Arial'
head_font = Font(name=F, bold=True, color='FFFFFF', size=10)
head_fill = PatternFill('solid', fgColor='1F3864')
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=14, bold=True)
note = Font(name=F, size=9, italic=True, color='595959')
input_fill = PatternFill('solid', fgColor='FFFF00')
wrap = Alignment(wrap_text=True, vertical='top')

def cond_pressure(status, i): return nodes[f'condition:{status}:{i}']['pressure']
def tag(side, p):
    word = 'Bearish' if p in ('B', 'HB') else 'Bullish' if p in ('U', 'HU') else 'Stable'
    return f"{side}-{word}{'(H)' if p in ('HB', 'HU') else ''}"
def pair(sp, rp): return nodes[f'pair:{sp}|{rp}']
def tug_text(sp, rp):
    return {'bears': 'Bears winning', 'bulls': 'Bulls winning', 'contested': 'Contested (sides pull apart)', 'none': 'Balanced'}[pair(sp, rp)['winner']]
def risky_text(sp, rp):
    n = pair(sp, rp)
    if n['veryRiskyAllowed']: return f"VERY RISKY {n['riskyDirection'].upper()} allowed at the nearest end of diversion (and a risky one at the tool's Risky level)"
    return f"Risky {n['riskyDirection'].upper()} allowed at the tool's Risky level" if n['riskyAllowed'] else 'No risky trade'
def trades(top, bottom, rare, sp, rp):
    if rare: return 'Do not trade (very volatile). For anyone who still trades: risky short at the top shown, risky long at the bottom shown.'
    out = [f'Short at {top} (resistance side)' if top else 'No short from a top (Bull Run)', f'Long at {bottom} (support side)' if bottom else 'No long from a bottom (Blood Bath)']
    if top in ('WTT-1',) or bottom in ('WTB+1',): out.append('WTT/WTB percentage is the key player for the weak-side level')
    return '; '.join(out)

wb = Workbook()
def sheet(name, headers, rows, widths, freeze='A2', filter_=True, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.append(headers)
    for c in ws[1]: c.font, c.fill, c.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical='center')
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font, c.alignment = body, wrap
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    if filter_ and rows: ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    return ws

# ---- All cases (529)
cases = []
for bi, box in enumerate(tool['boxes']):
    nr = len(box['resistanceConditions'])
    for n, row in enumerate(box['rows']):
        t, b, e = row.split('|')[:3]
        i, j = n // nr, n % nr
        rare = tool['explanations'][int(e)].startswith('RARE')
        top = None if t == 'BULL RUN' else t
        bottom = None if b == 'BLOOD BATH' else b
        sp, rp = cond_pressure(box['support'], i), cond_pressure(box['resistance'], j)
        outcome = 'Rare scenario: do not trade' if rare else ('Blood Bath (no bottom)' if not bottom else 'Bull Run (no top)' if not top else 'Top and bottom')
        label = f'{bi + 1}(S{i + 1}/R{j + 1})'
        cases.append([f'{bi + 1}|{i + 1}|{j + 1}', bi + 1, label, f"{label} {tag('S', sp)} {tag('R', rp)}", NAME[box['support']], NAME[box['resistance']], i + 1, box['supportConditions'][i], WORD[sp], j + 1, box['resistanceConditions'][j], WORD[rp], tug_text(sp, rp), top or 'none', bottom or 'none', outcome, trades(top, bottom, rare, sp, rp), risky_text(sp, rp), tool['explanations'][int(e)]])
cases_headers = ['Key (box|S|R)', 'COA box', 'Label', 'Signal tag', 'Support now', 'Resistance now', 'Support condition no.', 'Support history (tool wording)', 'Support pressure', 'Resistance condition no.', 'Resistance history (tool wording)', 'Resistance pressure', 'Tug of war', 'Expected TOP', 'Expected BOTTOM', 'Outcome', 'Trades possible', 'Risky trade (tug of war)', "The tool's explanation"]
col = {h: get_column_letter(k + 1) for k, h in enumerate(cases_headers)}

# ---- Read me
ws = wb.active; ws.title = 'Read me'
lines = [
    ('LTP scenarios: trading rules from the Detailed Chart of Accuracy', title),
    ('Built 2026-10-13 from the free course practice tool (529 cases captured and matched by our engine) and the knowledge graph.', note),
    ('', body),
    ('How to use', bold),
    ('1. Sheet "Lookup": type the COA box (1 to 9) and the support and resistance condition numbers in the yellow cells. The outcome fills in.', body),
    ('2. Sheet "All cases": every one of the 529 cases, filterable. The label reads scenario(S<support no.>/R<resistance no.>), for example 6(S3/R1); the signal tag adds the side tags: 6(S3/R1) S-Bearish R-Bullish(H).', body),
    ('3. Sheet "Conditions": the 23 numbered history conditions and the pressure each gives. Sheet "Pressure pairs": the pair of pressures and the top and bottom it gives.', body),
    ('4. Sheet "Scenarios": the nine boxes. Sheet "Rules": each rule with its source. Sheet "Open questions": what is still undecided.', body),
    ('', body),
    ('Legend', bold),
    ('Yellow cell = type here. Side tags: S-Bearish, R-Bullish, R-Stable, with (H) when high. Pressure: stable strong, bearish, bullish, highly bearish, highly bullish. Top and bottom: EOR / EOR+1 / WTT-1 and EOS / EOS-1 / WTB+1; "none" = no top (Bull Run) or no bottom (Blood Bath).', body),
    ('Source tags: course = free course or its practice tool; tool = observed on the live tool; owner = told to us by the system\'s owner; ours = our own choice.', body),
    ('', body),
    ('What this is not', bold),
    ('The practice tool gives the top and bottom only. Which entries are safe, moderate or risky in each scenario, and the stops, targets and times, are NOT in it and are still open (sheet "Open questions").', body),
    ('Tug of war (owner rule): a risky trade is allowed when both sides push the same way, or one pushes highly and the other is stable. Bears winning = risky short, bulls winning = risky long, at the level the live tool shows as Risky. Very high push (both the same way, at least one high): a very risky trade at the nearest end of diversion. Sides pulling against each other are contested: no risky trade. Decide from the outcome, not from the first-level scenario number.', body),
    ('Stop loss (owner rule): every trade risks exactly one diversion, one strike step from the entry: 50 points in NIFTY, 100 in BANKNIFTY and SENSEX, for every setup. The agent has it as v2.0 stop_rule one_diversion.', body),
    ('The pressure of each condition is our reading of the tool\'s wording, confirmed with the owner condition by condition on 2026-10-13. "~" in the app means the history is not one of the listed shapes.', body),
]
for r, (text, font) in enumerate(lines, 1):
    c = ws.cell(row=r, column=1, value=text); c.font = font; c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 140

# ---- Lookup (formulas into All cases)
lk = wb.create_sheet('Lookup')
n_cases = len(cases)
last = n_cases + 1
lk['A1'] = 'Lookup: pick a case, read the outcome'; lk['A1'].font = title
lk['A2'] = 'Type in the yellow cells. Box 1 to 9; condition numbers as in the tool (13 for a Strong side, 5 for a weak side; see sheet "Scenarios").'; lk['A2'].font = note
labels = [('COA box (1 to 9)', 4), ('Support condition no.', 5), ('Resistance condition no.', 6)]
for k, (text, row) in enumerate(labels):
    lk.cell(row=row, column=1, value=text).font = bold
    c = lk.cell(row=row, column=2, value=[6, 1, 1][k]); c.fill = input_fill; c.font = Font(name=F, size=10, bold=True, color='0000FF')
lk['C4'] = 'Example row: box 6, support condition 1, resistance condition 1 (both weak toward bottom, plain).'; lk['C4'].font = note
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='9', showErrorMessage=True, errorTitle='COA box', error='Enter a whole number from 1 to 9.')
dv2 = DataValidation(type='whole', operator='between', formula1='1', formula2='13', showErrorMessage=True, errorTitle='Condition number', error='Enter a whole number from 1 to 13 (a weak side has only 1 to 5).')
lk.add_data_validation(dv); lk.add_data_validation(dv2); dv.add('B4'); dv2.add('B5'); dv2.add('B6')
lk['A8'] = 'Key'; lk['B8'] = '=B4&"|"&B5&"|"&B6'
lk['A9'] = 'Row in "All cases"'; lk['B9'] = f"=IFERROR(MATCH(B8,'All cases'!$A$2:$A${last},0),0)"
out = [(h, col[h]) for h in ['Label', 'Signal tag', 'Support now', 'Resistance now', 'Support history (tool wording)', 'Support pressure', 'Resistance history (tool wording)', 'Resistance pressure', 'Tug of war', 'Expected TOP', 'Expected BOTTOM', 'Outcome', 'Trades possible', 'Risky trade (tug of war)', "The tool's explanation"]]
lk['A11'] = 'RESULT'; lk['A11'].font = bold
for k, (text, col) in enumerate(out):
    r = 12 + k
    lk.cell(row=r, column=1, value=text).font = bold
    c = lk.cell(row=r, column=2, value=f"=IF($B$9=0,\"Not a combination offered in this box\",INDEX('All cases'!${col}$2:${col}${last},$B$9))")
    c.font = body; c.alignment = wrap
for r in (8, 9): lk.cell(row=r, column=1).font = note; lk.cell(row=r, column=2).font = note
lk.column_dimensions['A'].width = 28; lk.column_dimensions['B'].width = 110; lk.column_dimensions['C'].width = 60

# ---- Scenarios
sc_rows = []
for bi, box in enumerate(tool['boxes']):
    ns, nr = len(box['supportConditions']), len(box['resistanceConditions'])
    sc_rows.append([bi + 1, NAME[box['support']], NAME[box['resistance']], nodes[f'scenario:{bi + 1}']['bias'], ns, nr, f'=E{bi + 2}*F{bi + 2}'])
ws = sheet('Scenarios', ['COA box (= scenario)', 'Support now', 'Resistance now', 'Basic bias', 'Support conditions', 'Resistance conditions', 'Cases'], sc_rows, [14, 26, 26, 40, 14, 14, 10], filter_=False)
ws.cell(row=len(sc_rows) + 2, column=6, value='Total').font = bold
tc = ws.cell(row=len(sc_rows) + 2, column=7, value=f'=SUM(G2:G{len(sc_rows) + 1})'); tc.font = bold

# ---- Conditions (23)
cond_rows = []
for status, label in (('strong', 'Strong side'), ('wtb', 'Weak toward bottom'), ('wtt', 'Weak toward top')):
    for n in sorted([x for x in graph['nodes'] if x['type'] == 'condition' and x['status'] == status], key=lambda x: x['number']):
        cond_rows.append([label, n['number'], n['wording'], ' > '.join(n['eventLog']), WORD[n['pressure']]])
sheet('Conditions', ['Applies to a side that is', 'Condition no.', "The tool's wording", 'Our event log (init, up, down, wtt, wtb, fromWtt, fromWtb)', 'Pressure it gives'], cond_rows, [24, 12, 100, 50, 18])

# ---- Pressure pairs (25)
pair_rows = []
for n in [x for x in graph['nodes'] if x['type'] == 'pair']:
    rare = n['rare']
    pair_rows.append([f"{tag('S', n['support'])} {tag('R', n['resistance'])}", WORD[n['support']], WORD[n['resistance']], tug_text(n['support'], n['resistance']), n['top'] or 'none', n['bottom'] or 'none', n.get('topWhenResistanceWeakTowardTop') or '', n.get('bottomWhenSupportWeakTowardBottom') or '', 'No (rare)' if rare else ('Yes' if n['tradable'] else 'No'), 'Rare scenario: do not trade' if rare else ('Blood Bath (no bottom)' if not n['bottom'] else 'Bull Run (no top)' if not n['top'] else 'Top and bottom'), trades(n['top'], n['bottom'], rare, n['support'], n['resistance']), risky_text(n['support'], n['resistance']), f"{n['riskyTop']} / {n['riskyBottom']}" if rare else ''])
sheet('Pressure pairs', ['Tags', 'Support pressure', 'Resistance pressure', 'Tug of war', 'Top', 'Bottom', 'Top if resistance is weak toward top and bullish', 'Bottom if support is weak toward bottom and bearish', 'Tradable', 'Outcome', 'Trades possible', 'Risky trade (tug of war)', 'Risky levels (rare scenario only)'], pair_rows, [28, 18, 18, 26, 10, 10, 22, 22, 12, 28, 60, 40, 18])

# ---- All cases
ws = sheet('All cases', cases_headers, cases, [14, 8, 12, 34, 22, 22, 10, 60, 16, 10, 60, 16, 26, 12, 12, 26, 60, 36, 70])
ws.freeze_panes = 'E2'

# ---- Rules & open questions
rule_rows = [[n['label'], n['source'], n['status'], n['detail']] for n in graph['nodes'] if n['type'] == 'rule']
sheet('Rules', ['Rule', 'Source', 'Status', 'Detail'], rule_rows, [60, 10, 12, 100])
q_rows = [[n['label']] for n in graph['nodes'] if n['type'] == 'open_question']
sheet('Open questions', ['Still undecided'], q_rows, [140], filter_=False)

wb.calculation.fullCalcOnLoad = True  # no spreadsheet engine here to fill cached values: Excel / Numbers / Sheets calculate on open
wb.save('docs/ltp-trading-rules.xlsx')
print('saved', n_cases, 'cases')
